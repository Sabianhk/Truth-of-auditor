"""Builder workbook export rekonsiliasi — dipakai export_run & export_center.

Satu jalur kode untuk sheet "Hasil" supaya export per-run dan per-batch
tidak pernah beda format.
"""
import re

from openpyxl import Workbook
from openpyxl.styles import Font

from web.templatetags.web_extras import reason_label

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Label relasi utk header kolom — di-inject dari views (hindari import melingkar).
_REL_LABELS_FALLBACK = ("Kiri", "Kanan")


def safe_name(s):
    """Nama file aman lintas OS: selain [A-Za-z0-9-] jadi '_'."""
    return re.sub(r"[^A-Za-z0-9-]+", "_", str(s or "")).strip("_")


# Awalan yang Excel/openpyxl baca sebagai formula (injeksi CSV/XLSX).
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def xlsx_safe(value):
    """Netralkan injeksi formula Excel pada sel string untrusted.

    String pihak ketiga (counterparty/username/description/nama file) yang
    berawalan =, +, -, @, tab, atau CR dibaca Excel sebagai formula —
    "=HYPERLINK(...)" di nama pengirim bank bisa dieksekusi saat auditor
    membuka laporan. Prefix apostrof membuat Excel menampilkannya sebagai
    teks literal. Nilai non-string (angka/tanggal/None) lolos apa adanya —
    kolom numerik/tanggal JANGAN dibungkus helper ini.
    """
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def batch_filename(batch):
    """rekonsiliasi_<toko>_<tanggal>.xlsx — permintaan UAT: tanggal + nama toko."""
    toko = safe_name(batch.toko.name if batch.toko else "toko")
    tgl = batch.recon_date.isoformat() if batch.recon_date else f"batch{batch.pk}"
    return f"rekonsiliasi_{toko}_{tgl}.xlsx"


def _sheet_title(base, existing):
    """Judul sheet <=31 char, tanpa karakter terlarang openpyxl, anti-duplikat."""
    t = re.sub(r"[\\/*?:\[\]]", "-", base)[:31]
    n, out = 2, t
    while out in existing:
        suffix = f" ({n})"
        out = t[: 31 - len(suffix)] + suffix
        n += 1
    existing.add(out)
    return out


def results_sheet(wb, run, title, rel_labels):
    """Tulis satu sheet 'Hasil' untuk `run` (kolom identik dgn export_run lama)."""
    L, R = rel_labels.get(run.relation, _REL_LABELS_FALLBACK)
    d = wb.create_sheet(title)
    headers = ["Status", f"{L} Ticket", f"{L} Nominal", f"{L} Username", f"{L} Nama Lengkap",
               f"{L} Player Bank", f"{L} Bank Title", f"{L} Handler", f"{L} Waktu",
               R, f"{R} Sumber", f"{R} Nominal", f"{R} Waktu", "Skor", "Alasan", "Detail"]
    d.append(headers)
    for c in d[1]:
        c.font = Font(bold=True)
    qs = run.results.select_related("left", "right", "left__source_type", "right__source_type")
    for r in qs.iterator():
        left, right = r.left, r.right
        # Sel string dari data pihak ketiga dibungkus xlsx_safe (anti injeksi
        # formula); kolom numerik/tanggal/label aplikasi tidak.
        d.append([
            r.get_bucket_display(),
            xlsx_safe(left.ticket_no) if left else "",
            float(left.amount) if left else "",
            xlsx_safe(left.username) if left else "",
            xlsx_safe(left.counterparty) if left else "",
            xlsx_safe((left.raw or {}).get("Player Bank", "")) if left else "",
            xlsx_safe((left.raw or {}).get("Bank Title", "")) if left else "",
            xlsx_safe((left.raw or {}).get("Handler", "")) if left else "",
            left.occurred_at.strftime("%d/%m %H:%M") if left and left.occurred_at else "",
            xlsx_safe(right.ticket_no or right.counterparty) if right else "",
            right.source_type.key if right else "",
            float(right.amount) if right else "",
            right.occurred_at.strftime("%d/%m %H:%M") if right and right.occurred_at else "",
            round(r.score or 0),
            reason_label(r.reason_code),
            xlsx_safe(r.reason_detail),
        ])
    return d


def build_batch_workbook(batch, batch_no, rel_labels):
    """Workbook satu batch: sheet Ringkasan + satu sheet Hasil per run."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    s = batch.summary or {}
    dp, wd, buckets = s.get("dp", {}), s.get("wd", {}), s.get("buckets", {})
    rows = [
        ("Toko", batch.toko.name if batch.toko else ""),
        ("Batch", f"#{batch_no}" if batch_no else f"#{batch.pk}"),
        ("Tanggal rekonsiliasi", batch.recon_date.strftime("%d/%m/%Y") if batch.recon_date else ""),
        ("Toleransi", f"{batch.tolerance.name} (±{batch.tolerance.date_window_days} hari)"),
        ("Dibuat", batch.created_at.strftime("%d/%m/%Y %H:%M")),
        ("", ""),
        ("DP Panel", dp.get("panel", 0)),
        ("DP Uang (matched)", dp.get("money_matched", dp.get("money", 0))),
        ("DP Selisih", dp.get("selisih", 0)),
        ("WD Panel", wd.get("panel", 0)),
        ("WD Uang (matched)", wd.get("money_matched", wd.get("money", 0))),
        ("WD Selisih", wd.get("selisih", 0)),
        ("", ""),
        ("Cocok", buckets.get("cocok", 0)),
        ("Perlu Ditinjau", buckets.get("perlu_tinjau", 0)),
        ("Tidak Cocok", buckets.get("tidak_cocok", 0)),
    ]
    for label, val in rows:
        ws.append([label, val])
    for cell in ws["A"]:
        cell.font = Font(bold=True)

    titles = {"Ringkasan"}
    for run in batch.runs.all().select_related("tolerance"):
        results_sheet(wb, run, _sheet_title(f"Hasil {run.get_relation_display()}", titles), rel_labels)
    return wb
