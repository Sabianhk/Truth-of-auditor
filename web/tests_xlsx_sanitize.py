"""W3-2: sanitasi injeksi formula Excel di semua jalur export xlsx.

String pihak ketiga (counterparty/username/description/reason_detail/nama file)
yang berawalan =, +, -, @ (atau tab/CR) dibaca Excel sebagai formula —
"=HYPERLINK(...)" di nama pengirim bank bisa mengeksekusi saat auditor membuka
laporan. Semua sel string untrusted dibungkus `xlsx_safe` (prefix apostrof).
"""
import io
from datetime import date, datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from reconciliation.models import MatchResult, MatchRun, ReconBatch, ToleranceProfile
from sources.models import SourceType, Toko, Upload
from transactions.models import Transaction
from web.exports import xlsx_safe

D = Decimal
PAYLOAD = '=HYPERLINK("http://evil.example","klik")'


def _no_formula_cells(ws):
    """Kumpulan nilai sel bertipe formula di satu sheet (harus kosong)."""
    return [c.value for row in ws.iter_rows() for c in row if c.data_type == "f"]


class XlsxSafeUnitTests(TestCase):
    def test_prefix_berbahaya_dinetralkan(self):
        for awal in ("=", "+", "-", "@", "\t"):
            s = awal + "PAYLOAD"
            self.assertEqual(xlsx_safe(s), "'" + s, s)

    def test_string_biasa_dan_non_string_lolos(self):
        self.assertEqual(xlsx_safe("BUDI SANTOSO"), "BUDI SANTOSO")
        self.assertEqual(xlsx_safe(""), "")
        self.assertEqual(xlsx_safe(None), None)
        self.assertEqual(xlsx_safe(50000.0), 50000.0)
        self.assertEqual(xlsx_safe(D("7")), D("7"))


class _Base(TestCase):
    def setUp(self):
        User = get_user_model()
        User.objects.create_user("aud", "a@a.co", "pw12345", role="supervisor")
        self.client.login(username="aud", password="pw12345")
        self.tol = ToleranceProfile.objects.get_or_create(
            name="Default", defaults={"date_window_days": 1}
        )[0]
        self.toko = Toko.objects.get(key="lbs")
        self.panel = SourceType.objects.get_or_create(key="panel", defaults={"name": "Panel"})[0]
        self.bank = SourceType.objects.get_or_create(key="bank", defaults={"name": "Bank"})[0]
        self.client.post(reverse("set_toko"), {"toko_id": self.toko.id})
        self._n = iter(range(1000))

    def tx(self, st, up, md, **kw):
        return Transaction.objects.create(
            upload=up, source_type=st, toko=self.toko, jenis=kw.pop("jenis", "depo"),
            amount=D(str(abs(md))), money_delta=D(str(md)),
            occurred_at=kw.pop("occurred_at", datetime(2026, 6, 27, 10, 0)),
            row_hash=f"xs{next(self._n)}", **kw,
        )

    def _wb(self, resp):
        self.assertEqual(resp.status_code, 200)
        return load_workbook(io.BytesIO(resp.content))


class ExportTransactionsSanitizeTests(_Base):
    def test_counterparty_formula_jadi_teks_literal(self):
        up = Upload.objects.create(source_type=self.bank, toko=self.toko)
        self.tx(self.bank, up, 10000, counterparty=PAYLOAD)
        r = self.client.get(reverse("transactions"), {"export": "1"})
        ws = self._wb(r).active
        self.assertEqual(_no_formula_cells(ws), [])
        vals = [c.value for row in ws.iter_rows(min_row=2) for c in row]
        self.assertIn("'" + PAYLOAD, vals)


class ExportRunSanitizeTests(_Base):
    def test_results_sheet_semua_kolom_untrusted_dibungkus(self):
        up = Upload.objects.create(source_type=self.panel, toko=self.toko)
        batch = ReconBatch.objects.create(toko=self.toko, tolerance=self.tol)
        run = MatchRun.objects.create(
            relation=MatchRun.Relation.PANEL_BANK, tolerance=self.tol, batch=batch)
        left = self.tx(self.panel, up, 10000, counterparty=PAYLOAD,
                       username="@budi", raw={"Player Bank": "=cmd|calc!A0"})
        MatchResult.objects.create(
            run=run, bucket=MatchResult.Bucket.TIDAK, reason_code="no_money",
            left=left, reason_detail="=EVIL()")
        r = self.client.get(reverse("export_run", args=[run.pk]))
        ws = self._wb(r)["Hasil"]
        self.assertEqual(_no_formula_cells(ws), [])
        vals = [c.value for row in ws.iter_rows(min_row=2) for c in row]
        self.assertIn("'" + PAYLOAD, vals)          # counterparty (Nama Lengkap)
        self.assertIn("'@budi", vals)               # username
        self.assertIn("'=cmd|calc!A0", vals)        # raw Player Bank
        self.assertIn("'=EVIL()", vals)             # reason_detail


class BatchUangSanitizeTests(_Base):
    def test_export_uang_tanpa_pasangan_dibungkus(self):
        from reconciliation.engine import run_batch

        up_b = Upload.objects.create(source_type=self.bank, toko=self.toko,
                                     original_name="=EVIL_FILE.csv")
        self.tx(self.bank, up_b, 14000, counterparty=PAYLOAD)
        batch = run_batch(self.toko, self.tol, recon_date=date(2026, 6, 27))
        r = self.client.get(reverse("batch_uang", args=[batch.pk]) + "?export=1")
        ws = self._wb(r).active
        self.assertEqual(_no_formula_cells(ws), [])
        vals = [c.value for row in ws.iter_rows(min_row=2) for c in row]
        self.assertIn("'" + PAYLOAD, vals)
        self.assertIn("'=EVIL_FILE.csv", vals)
